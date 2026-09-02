"""Builds the expense report spreadsheet.

Every figure here is computed in Python Decimal. When the AI phase arrives it
will write the narrative around these numbers -- it will never produce a
number itself, because a language model cannot be trusted with arithmetic.
"""
import uuid
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services.storage import user_dir

ZERO = Decimal("0.00")

# "#,##0.00" keeps trailing zeros, so 800 shows as 800.00 rather than 800.
MONEY_FORMAT = '#,##0.00'

_HEAD_FILL = PatternFill("solid", fgColor="1F4E5F")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, size=14)
_LABEL_FONT = Font(bold=True, size=10)
_TOTAL_FONT = Font(bold=True, size=11)
_THIN = Side(style="thin", color="C8D2D6")
_TOP_BORDER = Border(top=Side(style="medium", color="1F4E5F"))


def category_totals(expenses) -> list[tuple[str, Decimal, int]]:
    """Subtotals per category, largest first."""
    buckets: dict[str, list] = {}
    for e in expenses:
        entry = buckets.setdefault(e.category, [ZERO, 0])
        entry[0] += e.amount
        entry[1] += 1
    return sorted(
        ((cat, total, count) for cat, (total, count) in buckets.items()),
        key=lambda row: row[1],
        reverse=True,
    )


def grand_total(expenses) -> Decimal:
    """Summed in Decimal, never float, and never by a language model."""
    return sum((e.amount for e in expenses), ZERO)


def build_workbook(
    *,
    user,
    report_name: str,
    period_start: date,
    period_end: date,
    expenses,
    document_names: dict[int, str],
    currency: str = "INR",
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Expense Report"

    ws["A1"] = "EXPENSE REPORT"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = report_name
    ws["A2"].font = Font(size=11, color="4A5A60")

    for row, (label, value) in enumerate(
        [
            ("Employee", user.full_name),
            ("Employee ID", user.employee_id),
            ("Period", f"{period_start:%d %b %Y} to {period_end:%d %b %Y}"),
            ("Generated", f"{date.today():%d %b %Y}"),
        ],
        start=4,
    ):
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=row, column=2, value=value)

    headers = ["Date", "Category", "Paid to", "Description", "Receipt", f"Amount ({currency})"]
    header_row = 9
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = _HEAD_FONT
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="right" if col == len(headers) else "left")

    row = header_row + 1
    for e in expenses:
        ws.cell(row=row, column=1, value=e.expense_date).number_format = "DD MMM YYYY"
        ws.cell(row=row, column=2, value=e.category)
        ws.cell(row=row, column=3, value=e.vendor)
        ws.cell(row=row, column=4, value=e.description or "")
        ws.cell(row=row, column=5, value=document_names.get(e.document_id, ""))
        # Written as a float only because Excel has no decimal cell type. The
        # authoritative value stays in MySQL as DECIMAL; this is a rendering.
        amount = ws.cell(row=row, column=6, value=float(e.amount))
        amount.number_format = MONEY_FORMAT
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = Border(bottom=_THIN)
        row += 1

    row += 1
    ws.cell(row=row, column=5, value="Subtotals").font = _LABEL_FONT
    row += 1
    for cat, total, count in category_totals(expenses):
        ws.cell(row=row, column=5, value=f"{cat} ({count})")
        cell = ws.cell(row=row, column=6, value=float(total))
        cell.number_format = MONEY_FORMAT
        row += 1

    total_cell_row = row
    ws.cell(row=total_cell_row, column=5, value="TOTAL").font = _TOTAL_FONT
    total = ws.cell(row=total_cell_row, column=6, value=float(grand_total(expenses)))
    total.font = _TOTAL_FONT
    total.number_format = MONEY_FORMAT
    for col in (5, 6):
        ws.cell(row=total_cell_row, column=col).border = _TOP_BORDER

    ws.cell(row=total_cell_row + 2, column=1,
            value="All amounts calculated from verified expense records.").font = Font(
        size=9, italic=True, color="7A8A90")

    for col, width in enumerate([13, 14, 26, 30, 24, 15], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    # Keeps the column headers visible when scrolling a long report.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # A UUID rather than a timestamp: two reports for the same period would
    # otherwise share a filename, and the second would overwrite the first's
    # file while both rows still pointed at it.
    filename = f"report_{period_start:%Y%m%d}_{period_end:%Y%m%d}_{uuid.uuid4().hex[:8]}.xlsx"
    path = user_dir(user.user_id, "reports") / filename
    wb.save(path)
    return path
